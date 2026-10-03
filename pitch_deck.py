from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from io import BytesIO


# ============================================================
# COLORS
# ============================================================

NAVY = RGBColor(20, 35, 60)
BLUE = RGBColor(45, 100, 180)
LIGHT_BLUE = RGBColor(225, 238, 250)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(35, 35, 35)
GREY = RGBColor(100, 100, 100)
LIGHT_GREY = RGBColor(242, 244, 247)
GREEN = RGBColor(50, 150, 100)


# ============================================================
# HELPERS
# ============================================================

def set_background(slide, color=WHITE):

    fill = slide.background.fill

    fill.solid()

    fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    font_size=24,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT
):

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    tf = box.text_frame

    tf.clear()

    p = tf.paragraphs[0]

    p.alignment = align

    run = p.add_run()

    run.text = str(text)

    run.font.size = Pt(font_size)

    run.font.bold = bold

    run.font.color.rgb = color

    return box


def add_header(slide, title, subtitle=None):

    add_text(
        slide,
        title,
        0.7,
        0.45,
        11.5,
        0.7,
        font_size=28,
        color=NAVY,
        bold=True
    )

    if subtitle:

        add_text(
            slide,
            subtitle,
            0.7,
            1.1,
            11.2,
            0.5,
            font_size=13,
            color=GREY
        )


def add_footer(slide, number):

    add_text(
        slide,
        f"{number} / 10",
        11.4,
        7.05,
        1.0,
        0.25,
        font_size=9,
        color=GREY,
        align=PP_ALIGN.RIGHT
    )


def add_card(
    slide,
    title,
    body,
    x,
    y,
    w,
    h
):

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    shape.fill.solid()

    shape.fill.fore_color.rgb = LIGHT_BLUE

    shape.line.color.rgb = LIGHT_BLUE

    add_text(
        slide,
        title,
        x + 0.25,
        y + 0.2,
        w - 0.5,
        0.45,
        font_size=17,
        color=NAVY,
        bold=True
    )

    add_text(
        slide,
        body,
        x + 0.25,
        y + 0.75,
        w - 0.5,
        h - 0.9,
        font_size=13,
        color=DARK
    )


# ============================================================
# CREATE PITCH DECK
# ============================================================

def create_pitch_deck(
    profile,
    startup,
    portfolio
):

    prs = Presentation()

    prs.slide_width = Inches(13.333)

    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]


    # ========================================================
    # SLIDE 1 — COVER
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide, NAVY)

    add_text(
        slide,
        startup.get(
            "startup_name",
            "Startup Name"
        ),
        0.8,
        1.65,
        11.7,
        1.0,
        font_size=42,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        startup.get(
            "industry",
            "Industry"
        ),
        0.8,
        2.7,
        11.7,
        0.6,
        font_size=20,
        color=LIGHT_BLUE,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        startup.get(
            "solution",
            "Innovative solution"
        ),
        1.3,
        3.55,
        10.7,
        1.0,
        font_size=18,
        color=WHITE,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        profile.get(
            "name",
            "Founder"
        ),
        1,
        5.8,
        11.3,
        0.5,
        font_size=16,
        color=LIGHT_BLUE,
        align=PP_ALIGN.CENTER
    )


    # ========================================================
    # SLIDE 2 — PROBLEM
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "The Problem",
        "The customer challenge we are solving"
    )

    add_card(
        slide,
        "Customer Pain Point",
        startup.get(
            "problem",
            "Problem not provided."
        ),
        0.8,
        1.8,
        11.7,
        2.0
    )

    add_card(
        slide,
        "Target Market",
        startup.get(
            "target_market",
            "Target market not provided."
        ),
        0.8,
        4.2,
        5.6,
        1.6
    )

    add_card(
        slide,
        "Market Opportunity",
        startup.get(
            "market_size",
            "Market size not provided."
        ),
        6.9,
        4.2,
        5.6,
        1.6
    )

    add_footer(slide, 2)


    # ========================================================
    # SLIDE 3 — SOLUTION
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "Our Solution",
        "How the startup addresses the problem"
    )

    add_card(
        slide,
        "Solution",
        startup.get(
            "solution",
            "Solution not provided."
        ),
        0.8,
        1.8,
        7.3,
        3.8
    )

    add_card(
        slide,
        "Key Advantage",
        startup.get(
            "competitive_advantage",
            "Competitive advantage not provided."
        ),
        8.4,
        1.8,
        4.1,
        3.8
    )

    add_footer(slide, 3)


    # ========================================================
    # SLIDE 4 — MARKET
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "Market Opportunity",
        "Target market and growth potential"
    )

    add_card(
        slide,
        "Target Market",
        startup.get(
            "target_market",
            "Not provided."
        ),
        0.8,
        1.7,
        5.6,
        2.0
    )

    add_card(
        slide,
        "Market Size",
        startup.get(
            "market_size",
            "Not provided."
        ),
        6.9,
        1.7,
        5.6,
        2.0
    )

    add_card(
        slide,
        "Expected Growth",
        startup.get(
            "growth_rate",
            "Not provided."
        ),
        0.8,
        4.2,
        5.6,
        1.7
    )

    add_card(
        slide,
        "Industry",
        startup.get(
            "industry",
            "Not provided."
        ),
        6.9,
        4.2,
        5.6,
        1.7
    )

    add_footer(slide, 4)


    # ========================================================
    # SLIDE 5 — BUSINESS MODEL
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "Business Model",
        "How the startup creates and captures value"
    )

    add_card(
        slide,
        "Revenue Model",
        startup.get(
            "revenue_model",
            "Revenue model not provided."
        ),
        0.8,
        1.8,
        7.0,
        2.6
    )

    add_card(
        slide,
        "Current Revenue",
        startup.get(
            "current_revenue",
            "Not provided."
        ),
        8.2,
        1.8,
        4.3,
        2.6
    )

    add_card(
        slide,
        "Funding Stage",
        startup.get(
            "funding_stage",
            "Not provided."
        ),
        0.8,
        4.8,
        11.7,
        1.4
    )

    add_footer(slide, 5)


    # ========================================================
    # SLIDE 6 — COMPETITIVE ADVANTAGE
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "Competitive Advantage",
        "Why the startup can differentiate"
    )

    add_card(
        slide,
        "Differentiation",
        startup.get(
            "competitive_advantage",
            "Not provided."
        ),
        0.8,
        1.8,
        11.7,
        2.3
    )

    add_card(
        slide,
        "Customer Problem",
        startup.get(
            "problem",
            "Not provided."
        ),
        0.8,
        4.5,
        5.6,
        1.5
    )

    add_card(
        slide,
        "Solution",
        startup.get(
            "solution",
            "Not provided."
        ),
        6.9,
        4.5,
        5.6,
        1.5
    )

    add_footer(slide, 6)


    # ========================================================
    # SLIDE 7 — TRACTION
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "Traction & Financial Snapshot",
        "Current progress and financial indicators"
    )

    add_card(
        slide,
        "Current Revenue",
        startup.get(
            "current_revenue",
            "Not provided."
        ),
        0.8,
        1.7,
        3.6,
        2.0
    )

    add_card(
        slide,
        "Growth Rate",
        startup.get(
            "growth_rate",
            "Not provided."
        ),
        4.9,
        1.7,
        3.6,
        2.0
    )

    add_card(
        slide,
        "Funding Stage",
        startup.get(
            "funding_stage",
            "Not provided."
        ),
        9.0,
        1.7,
        3.5,
        2.0
    )

    add_card(
        slide,
        "Portfolio Achievement",
        portfolio.get(
            "project_result",
            "Portfolio result not provided."
        ),
        0.8,
        4.4,
        11.7,
        1.7
    )

    add_footer(slide, 7)


    # ========================================================
    # SLIDE 8 — FOUNDER
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "Founder & Team",
        "The people behind the startup"
    )

    add_card(
        slide,
        profile.get(
            "name",
            "Founder"
        ),
        profile.get(
            "bio",
            "Founder information not provided."
        ),
        0.8,
        1.8,
        7.0,
        3.5
    )

    add_card(
        slide,
        "Skills",
        profile.get(
            "skills",
            "Skills not provided."
        ),
        8.2,
        1.8,
        4.3,
        1.5
    )

    add_card(
        slide,
        "Role",
        profile.get(
            "role",
            "Role not provided."
        ),
        8.2,
        3.8,
        4.3,
        1.5
    )

    add_footer(slide, 8)


    # ========================================================
    # SLIDE 9 — ROADMAP
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide)

    add_header(
        slide,
        "Future Roadmap",
        "Next steps for growth and development"
    )

    add_card(
        slide,
        "Roadmap",
        portfolio.get(
            "roadmap",
            "Roadmap not provided."
        ),
        0.8,
        1.8,
        11.7,
        3.5
    )

    add_card(
        slide,
        "Major Achievement",
        portfolio.get(
            "achievements",
            "Achievements not provided."
        ),
        0.8,
        5.6,
        11.7,
        1.0
    )

    add_footer(slide, 9)


    # ========================================================
    # SLIDE 10 — CLOSING
    # ========================================================

    slide = prs.slides.add_slide(blank_layout)

    set_background(slide, NAVY)

    add_text(
        slide,
        "Let's Build the Future",
        0.8,
        1.7,
        11.7,
        1.0,
        font_size=38,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        startup.get(
            "startup_name",
            "Startup"
        ),
        0.8,
        3.0,
        11.7,
        0.7,
        font_size=25,
        color=LIGHT_BLUE,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        startup.get(
            "solution",
            "Innovative solution"
        ),
        1.5,
        4.0,
        10.3,
        1.0,
        font_size=17,
        color=WHITE,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        profile.get(
            "email",
            ""
        ),
        1,
        5.8,
        11.3,
        0.5,
        font_size=15,
        color=LIGHT_BLUE,
        align=PP_ALIGN.CENTER
    )


    # ========================================================
    # RETURN POWERPOINT
    # ========================================================

    output = BytesIO()

    prs.save(output)

    output.seek(0)

    return output